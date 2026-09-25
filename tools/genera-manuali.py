"""Genera i manuali Word (M-BPM-*.docx) dalla documentazione, usando il modello aziendale.

Uso:
    python tools/genera-manuali.py integrazione      # un manuale
    python tools/genera-manuali.py                   # tutti i manuali

Per ogni manuale:
1. compone il Markdown delle sezioni indicate (stessa logica di tools/export-docx.py);
2. lo converte in Word con Pandoc, usando il modello come riferimento per gli stili;
3. parte dal modello (export/M-BPM-template.docx), ne tiene la copertina e l'indice,
   sostituisce il titolo, toglie i contenuti di esempio e accoda il contenuto generato.

L'indice si aggiorna a mano in Word prima di produrre il PDF.
"""
import importlib.util
import re
import shutil
import zipfile
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn
from docxcompose.composer import Composer

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
EXPORT_DIR = ROOT / "export"
TEMPLATE = EXPORT_DIR / "M-BPM-template.docx"

MANUALI = {
    "designer": {
        "file": "M-BPM-Designer.docx",
        "titolo": "Manuale Designer BPM",
        "sezioni": ["modelli-di-processo", "classi-documentali", "parti-comuni"],
    },
    "integrazione": {
        "file": "M-BPM-Integrazione.docx",
        "titolo": "Manuale Integrazione BPM",
        "sezioni": ["integrazione"],
    },
    "ai": {
        "file": "M-BPM-AI.docx",
        "titolo": "Manuale Intelligenza artificiale BPM",
        "sezioni": ["intelligenza-artificiale"],
    },
    "admin": {
        "file": "M-BPM-Admin.docx",
        "titolo": "Manuale Amministrazione BPM",
        "sezioni": ["amministrazione"],
    },
}

# Stili generati da Pandoc da ricondurre agli stili del modello aziendale (per styleId).
MAPPA_STILI_PARAGRAFO = {
    "BodyText": "TestomanualeCSW1",
    "FirstParagraph": "TestomanualeCSW1",
    "Compact": "TestomanualeCSW1",
}
STILE_TABELLA = "TableGrid"
COLORE_BORDI_TABELLA = "808080"   # grigio scuro (Shape: testo secondario)
SPAZIO_ATTORNO = "160"            # 8 pt prima e dopo tabelle e riquadri

# Ordine degli elementi di w:pPr e w:tblPr secondo lo schema (quelli che seguono gli elementi inseriti).
DOPO_SPACING = ["ind", "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc", "textDirection",
                "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr", "pPrChange"]
DOPO_TBLBORDERS = ["shd", "tblLayout", "tblCellMar", "tblLook", "tblCaption", "tblDescription", "tblPrChange"]


def inserisci_in_ordine(padre, elemento, successivi: list[str]) -> None:
    tag_successivi = {qn(f"w:{n}") for n in successivi}
    dopo = next((el for el in padre if el.tag in tag_successivi), None)
    if dopo is not None:
        dopo.addprevious(elemento)
    else:
        padre.append(elemento)


def imposta_spaziatura(p, before: str | None = None, after: str | None = None) -> None:
    p_pr = p.find(qn("w:pPr"))
    if p_pr is None:
        p_pr = p.makeelement(qn("w:pPr"), {})
        p.insert(0, p_pr)
    spacing = p_pr.find(qn("w:spacing"))
    if spacing is None:
        spacing = p_pr.makeelement(qn("w:spacing"), {})
        inserisci_in_ordine(p_pr, spacing, DOPO_SPACING)
    if before is not None:
        spacing.set(qn("w:before"), before)
    if after is not None:
        spacing.set(qn("w:after"), after)


def stile_paragrafo(el) -> str:
    return el.xpath("string(./w:pPr/w:pStyle/@w:val)") if el.tag == qn("w:p") else ""

# Link rimasti relativi dopo la conversione: puntano a pagine fuori dal manuale o a file del sito.
LINK_RELATIVO_RE = re.compile(r"(?<!!)\[([^\]]*)\]\(([^)\s]+)\)(\{[^}]*\})?")


def sistema_link_esterni(testo: str) -> str:
    """Link a pagine di altre sezioni: resta il testo. Link a file: rimando al sito."""
    def sostituisci(m: re.Match) -> str:
        etichetta, destinazione = m.group(1), m.group(2)
        if destinazione.startswith(("#", "http://", "https://", "mailto:")):
            return m.group(0)
        percorso = destinazione.split("#", 1)[0]
        if percorso.endswith(".md") or percorso.endswith("/") or percorso == "":
            return etichetta
        return f"{etichetta} (disponibile sul sito della documentazione)"
    return LINK_RELATIVO_RE.sub(sostituisci, testo)


def carica_export():
    spec = importlib.util.spec_from_file_location("export_docx", ROOT / "tools" / "export-docx.py")
    modulo = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(modulo)
    return modulo


def titolo_sezione(ex, sezione: str) -> str:
    for voce in ex.load_nav(DOCS / ".nav.yml"):
        if isinstance(voce, dict):
            for titolo, valore in voce.items():
                if valore == sezione:
                    return titolo
    return sezione.replace("-", " ").title()


def componi_markdown(ex, sezioni: list[str], destinazione: Path) -> int:
    pages = []
    for sezione in sezioni:
        nav = DOCS / sezione / ".nav.yml"
        sezione_pages = ex.walk_nav(nav, depth=1, title_override=titolo_sezione(ex, sezione))
        if len(sezioni) == 1:
            # Manuale di una sola sezione: le pagine principali diventano argomenti di livello 1.
            for p in sezione_pages[1:]:
                p.depth = max(1, p.depth - 1)
        pages.extend(sezione_pages)

    rendered, headings = {}, {}
    for page in pages:
        testo, titoli = ex.render_page(page)
        rendered[page.path], headings[page.path] = testo, titoli

    destinazione.write_text("\n\n".join(rendered[p.path] for p in pages) + "\n", encoding="utf-8")
    pandoc = ex.find_pandoc()
    native = ex.run_pandoc_native(pandoc, destinazione)
    ids = ex.HEADER_ID_RE.findall(native)
    page_top_id, fragment_map = ex.build_id_maps(pages, headings, ids)
    finale = "\n\n".join(
        ex.rewrite_internal_links(rendered[p.path], p.path.parent, page_top_id, fragment_map) for p in pages
    )
    destinazione.write_text(sistema_link_esterni(finale) + "\n", encoding="utf-8")
    return len(pages)


def converti_con_pandoc(ex, markdown: Path, uscita: Path, riferimento: Path) -> None:
    subprocess.run(
        [
            ex.find_pandoc(),
            str(markdown),
            "--from=markdown+raw_html",
            "--to=docx",
            f"--reference-doc={riferimento}",
            f"--resource-path={DOCS}",
            "-o",
            str(uscita),
        ],
        check=True,
    )


def stili_mancanti(ex) -> str:
    """Stili del riferimento predefinito di Pandoc assenti nel modello (Table, Compact, Source Code…).

    Senza di essi alcuni programmi (per esempio LibreOffice) non impaginano correttamente le tabelle.
    """
    with tempfile.TemporaryDirectory() as tmp:
        ref = Path(tmp) / "reference.docx"
        with open(ref, "wb") as out:
            subprocess.run([ex.find_pandoc(), "--print-default-data-file", "reference.docx"], stdout=out, check=True)
        with zipfile.ZipFile(ref) as z:
            predefiniti = z.read("word/styles.xml").decode("utf-8")
    with zipfile.ZipFile(TEMPLATE) as z:
        modello = z.read("word/styles.xml").decode("utf-8")
    presenti = set(re.findall(r'w:styleId="([^"]+)"', modello))
    aggiunte = []
    for m in re.finditer(r'<w:style [^>]*w:styleId="([^"]+)".*?</w:style>', predefiniti, re.S):
        if m.group(1) not in presenti:
            aggiunte.append(re.sub(r'\s*w:default="1"', "", m.group(0)))
    testo = "".join(aggiunte)
    return testo


def crea_riferimento(ex, destinazione: Path, aggiunte: str) -> None:
    """Copia del modello con gli stili mancanti, da usare come --reference-doc di Pandoc."""
    with zipfile.ZipFile(TEMPLATE) as zin, zipfile.ZipFile(destinazione, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            dati = zin.read(item.filename)
            if item.filename == "word/styles.xml":
                dati = dati.decode("utf-8").replace("</w:styles>", aggiunte + "</w:styles>").encode("utf-8")
            zout.writestr(item, dati)


def prepara_modello(titolo: str, riferimento: Path) -> Document:
    doc = Document(riferimento)
    corpo = doc.element.body
    figli = list(corpo.iterchildren())

    # Titolo in copertina.
    for p in doc.paragraphs:
        if p.style.name == "CSW TITOLO":
            p.runs[0].text = titolo
            for run in p.runs[1:]:
                run.text = ""
            break

    # Tiene copertina e indice; toglie i contenuti di esempio che seguono l'indice.
    ultimo_toc = max(i for i, el in enumerate(figli) if el.tag == qn("w:p") and "TOC" in (el.xpath("string(.//w:pStyle/@w:val)") or ""))
    inizio_esempi = next(
        i for i, el in enumerate(figli)
        if i > ultimo_toc and el.tag == qn("w:p") and el.xpath("string(.//w:pStyle/@w:val)") == "Heading1"
    )
    for el in figli[inizio_esempi:]:
        if el.tag != qn("w:sectPr"):
            corpo.remove(el)

    # Salto pagina dopo l'indice.
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # Proprietà del documento. L'indice si aggiorna a mano in Word (clic destro > Aggiorna campo):
    # non si imposta w:updateFields, che farebbe comparire un avviso a ogni apertura.
    doc.core_properties.title = titolo
    return doc


def imposta_stili_codice(doc: Document) -> None:
    """Codice più piccolo del testo; blocchi di codice su fondo grigio chiaro (colori Shape)."""
    stili = doc.styles.element
    for stile in stili.iter(qn("w:style")):
        sid = stile.get(qn("w:styleId"))
        if sid not in ("VerbatimChar", "SourceCode"):
            continue
        r_pr = stile.find(qn("w:rPr"))
        if r_pr is None:
            r_pr = stile.makeelement(qn("w:rPr"), {})
            stile.append(r_pr)
        for tag in ("w:sz", "w:szCs"):
            for el in r_pr.findall(qn(tag)):
                r_pr.remove(el)
            r_pr.append(r_pr.makeelement(qn(tag), {qn("w:val"): "18"}))
        if sid == "SourceCode":
            p_pr = stile.find(qn("w:pPr"))
            if p_pr is None:
                p_pr = stile.makeelement(qn("w:pPr"), {})
                r_pr.addprevious(p_pr)
            if p_pr.find(qn("w:shd")) is None:
                shd = p_pr.makeelement(qn("w:shd"), {qn("w:val"): "clear", qn("w:color"): "auto", qn("w:fill"): "F4F6F9"})
                ww = p_pr.find(qn("w:wordWrap"))
                (ww.addprevious(shd) if ww is not None else p_pr.append(shd))


def applica_stili(doc: Document) -> None:
    imposta_stili_codice(doc)
    corpo = doc.element.body
    # Riquadri (note, avvertenze): testo del manuale rientrato, con filo a sinistra.
    for p_style in list(corpo.iter(qn("w:pStyle"))):
        if p_style.get(qn("w:val")) == "BlockText":
            p_style.set(qn("w:val"), "TestomanualeCSW1")
            p_pr = p_style.getparent()
            ind = p_pr.makeelement(qn("w:ind"), {qn("w:left"): "567"})
            bdr = p_pr.makeelement(qn("w:pBdr"), {})
            left = bdr.makeelement(qn("w:left"), {qn("w:val"): "single", qn("w:sz"): "12", qn("w:space"): "8", qn("w:color"): "3093EF"})
            bdr.append(left)
            p_pr.append(bdr)
            p_pr.append(ind)
            p_style.set("riquadro", "1")
    for p_style in corpo.iter(qn("w:pStyle")):
        nuovo = MAPPA_STILI_PARAGRAFO.get(p_style.get(qn("w:val")))
        if nuovo:
            p_style.set(qn("w:val"), nuovo)
    # Ordine degli elementi richiesto dallo schema: pStyle sempre per primo in pPr,
    # abstractNumId sempre per primo in w:num.
    for p_pr in corpo.iter(qn("w:pPr")):
        p_style = p_pr.find(qn("w:pStyle"))
        if p_style is not None and p_pr.index(p_style) != 0:
            p_pr.remove(p_style)
            p_pr.insert(0, p_style)
        bdr, ind = p_pr.find(qn("w:pBdr")), p_pr.find(qn("w:ind"))
        for el in (bdr, ind):
            if el is not None:
                p_pr.remove(el)
        # pBdr dopo pStyle/keepNext/…/numPr; ind prima di jc: posizioni sicure.
        if bdr is not None:
            num = p_pr.find(qn("w:numPr"))
            (num if num is not None else p_pr[0]).addnext(bdr)
        if ind is not None:
            jc = p_pr.find(qn("w:jc"))
            if jc is not None:
                jc.addprevious(ind)
            else:
                rpr = p_pr.find(qn("w:rPr"))
                rpr.addprevious(ind) if rpr is not None else p_pr.append(ind)
    try:
        numerazione = doc.part.numbering_part.element
        for num in numerazione.iter(qn("w:num")):
            abstract = num.find(qn("w:abstractNumId"))
            if abstract is not None and num.index(abstract) != 0:
                num.remove(abstract)
                num.insert(0, abstract)
            # Al massimo un lvlOverride per livello (l'unione dei documenti può duplicarli).
            visti = {}
            for lo in num.findall(qn("w:lvlOverride")):
                livello = lo.get(qn("w:ilvl"))
                if livello in visti:
                    num.remove(visti[livello])
                visti[livello] = lo
    except NotImplementedError:
        pass
    # Spazio prima e dopo i riquadri (sequenze di paragrafi con filo a sinistra).
    figli = list(corpo)
    riquadro = lambda el: el.tag == qn("w:p") and el.find("./w:pPr/w:pStyle[@riquadro='1']", {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}) is not None
    for i, el in enumerate(figli):
        if riquadro(el):
            if i == 0 or not riquadro(figli[i - 1]):
                imposta_spaziatura(el, before=SPAZIO_ATTORNO)
            if i + 1 == len(figli) or not riquadro(figli[i + 1]):
                imposta_spaziatura(el, after=SPAZIO_ATTORNO)
    for el in corpo.iter(qn("w:pStyle")):
        el.attrib.pop("riquadro", None)

    # Spazio dopo le tabelle: il paragrafo che segue (se non è un titolo) si stacca un po'.
    for i, el in enumerate(figli):
        if el.tag == qn("w:tbl") and i + 1 < len(figli):
            succ = figli[i + 1]
            if succ.tag == qn("w:p") and not stile_paragrafo(succ).startswith("Heading"):
                imposta_spaziatura(succ, before=SPAZIO_ATTORNO)

    for tbl in corpo.iter(qn("w:tbl")):
        tbl_pr = tbl.find(qn("w:tblPr"))
        # Bordi grigio scuro invece che neri.
        vecchi = tbl_pr.find(qn("w:tblBorders"))
        if vecchi is not None:
            tbl_pr.remove(vecchi)
        bordi = tbl_pr.makeelement(qn("w:tblBorders"), {})
        for lato in ("top", "left", "bottom", "right", "insideH", "insideV"):
            bordi.append(bordi.makeelement(qn(f"w:{lato}"), {
                qn("w:val"): "single", qn("w:sz"): "4", qn("w:space"): "0", qn("w:color"): COLORE_BORDI_TABELLA}))
        inserisci_in_ordine(tbl_pr, bordi, DOPO_TBLBORDERS)
        for tc_bordi in tbl.iter(qn("w:tcBorders")):
            for lato in tc_bordi:
                lato.set(qn("w:color"), COLORE_BORDI_TABELLA)
                if lato.get(qn("w:sz")) is None:
                    lato.set(qn("w:sz"), "4")
        stile = tbl_pr.find(qn("w:tblStyle"))
        if stile is None:
            stile = tbl_pr.makeelement(qn("w:tblStyle"), {})
            tbl_pr.insert(0, stile)
        stile.set(qn("w:val"), STILE_TABELLA)
        # Tabelle a tutta larghezza. Alcune versioni di Pandoc scrivono larghezze
        # decimali (es. "5000.0") che non tutti i programmi interpretano.
        tbl_w = tbl_pr.find(qn("w:tblW"))
        if tbl_w is None:
            tbl_w = tbl_pr.makeelement(qn("w:tblW"), {})
            tbl_pr.append(tbl_w)
        tbl_w.set(qn("w:type"), "pct")
        tbl_w.set(qn("w:w"), "5000")


def genera(chiave: str) -> Path:
    ex = carica_export()
    conf = MANUALI[chiave]
    uscita = EXPORT_DIR / conf["file"]
    with tempfile.TemporaryDirectory() as tmp:
        md = Path(tmp) / "manuale.md"
        corpo = Path(tmp) / "corpo.docx"
        riferimento = Path(tmp) / "riferimento.docx"
        crea_riferimento(ex, riferimento, stili_mancanti(ex))
        n = componi_markdown(ex, conf["sezioni"], md)
        converti_con_pandoc(ex, md, corpo, riferimento)
        doc = prepara_modello(conf["titolo"], riferimento)
        composer = Composer(doc)
        composer.append(Document(corpo))
        applica_stili(composer.doc)
        composer.save(uscita)
    print(f"{uscita.relative_to(ROOT)}: {n} pagine")
    return uscita


def main() -> int:
    chiavi = sys.argv[1:] or list(MANUALI)
    for chiave in chiavi:
        if chiave not in MANUALI:
            print(f"Manuale sconosciuto: {chiave}. Disponibili: {', '.join(MANUALI)}")
            return 2
        genera(chiave)
    return 0


if __name__ == "__main__":
    sys.exit(main())
